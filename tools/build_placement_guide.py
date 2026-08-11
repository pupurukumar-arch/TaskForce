from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/purukumar/Desktop/TaskForce_Orbit_Placement_Interview_Guide.docx"
BLUE = "1F4D78"; LIGHT_BLUE = "E8EEF5"; INK = "172033"; MUTED = "52606D"; GOLD = "7A5A00"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def set_cell_width(cell, inches):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn("w:tcW"))
    if tcW is None: tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(inches * 1440))); tcW.set(qn("w:type"), "dxa")

def font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"; run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic; run.font.color.rgb = RGBColor.from_string(color)

def add_text(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        font(p.add_run(bold_prefix), bold=True)
        font(p.add_run(text[len(bold_prefix):]))
    else: font(p.add_run(text))
    return p

def h(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); font(r, {1:16,2:13,3:12}[level], bold=True, color=BLUE if level < 3 else "1F4D78")
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.15; font(p.add_run(text)); return p

NUMBERING = {}

def numbered(doc, text, sequence):
    if sequence not in NUMBERING:
        root = doc.part.numbering_part.element
        style_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
        base_num = next(item for item in root.findall(qn("w:num")) if int(item.get(qn("w:numId"))) == style_num_id)
        abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
        next_id = max([int(item.get(qn("w:numId"))) for item in root.findall(qn("w:num"))] + [0]) + 1
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(next_id))
        abstract = OxmlElement("w:abstractNumId"); abstract.set(qn("w:val"), abstract_id); num.append(abstract)
        override = OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride"); start.set(qn("w:val"), "1"); override.append(start); num.append(override); root.append(num)
        NUMBERING[sequence] = next_id
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr(); numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); numPr.append(ilvl)
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), str(NUMBERING[sequence])); numPr.append(numId); pPr.append(numPr)
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.15; font(p.add_run(text)); return p

def callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.autofit = False
    cell = table.cell(0,0); set_cell_width(cell, 6.5); set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); font(p.add_run(title), 10.5, True, BLUE)
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(3); font(p2.add_run(body), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False; t.style = "Table Grid"
    for i, label in enumerate(headers):
        cell=t.rows[0].cells[i]; set_cell_width(cell,widths[i]); set_cell_shading(cell,LIGHT_BLUE); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); font(p.add_run(label),9.5,True,BLUE)
    for row in rows:
        cells=t.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i],widths[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.05; font(p.add_run(value),9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    return t

def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Task Force Orbit | "); font(run, 8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); paragraph._p.append(fld)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(0.8); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.85); sec.right_margin=Inches(0.85)
sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
styles=doc.styles
styles["Normal"].font.name="Calibri"; styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); styles["Normal"].font.size=Pt(11)
for name in ["Heading 1","Heading 2","Heading 3"]:
    styles[name].font.name="Calibri"; styles[name].font.color.rgb=RGBColor.from_string(BLUE)
footer=sec.footer.paragraphs[0]; page_number(footer)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.LEFT; font(header.add_run("PLACEMENT INTERVIEW REFERENCE GUIDE"),8.5,True,MUTED)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(28)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(7); font(p.add_run("TASK FORCE ORBIT"),28,True,BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(24); font(p.add_run("End-to-End Placement Interview Guide"),15,False,MUTED)
callout(doc,"How to use this guide","Read the 60-second pitch first. Then revise the architecture, authentication, workflow, deployment, and the question bank. Answer with the facts of this codebase; do not claim features that are marked as future work.")
table(doc,["Project","Audience","Goal"],[["Task Force Orbit","Final-year B.Tech placement interview","Explain the product, implementation decisions, security, deployment, and trade-offs with confidence."]],[1.25,2.1,3.15])
add_text(doc,"Prepared from the current TaskForce repository and focused web research. It is intentionally practical: enough depth for a fresher interview, without pretending to be a production-scale enterprise system.",italic:=None)
doc.add_page_break()

h(doc,"1. The 60-Second Project Pitch")
add_text(doc,"Task Force Orbit is a role-based project collaboration web application. It solves a practical gap in simple task boards: work should not only be created and marked done; a member should be able to attach evidence, submit the work for review, and an authorized project leader should approve it or send it back for changes.")
callout(doc,"Say this in an interview","I built a full-stack team-workflow application using React, Express, MongoDB, JWT authentication, private S3 attachments, SMTP email flows, and Vercel deployment. The strongest part is the backend authorization: the UI may hide buttons, but every sensitive API route independently checks the authenticated user and project role.")
h(doc,"What the product does",2)
for x in ["Email-verified registration, login, logout, password reset, and refresh-token session renewal.","Projects with Admin, Project Admin, and Member roles; invitations for users who have not registered yet.","Tasks with assignee, priority, difficulty, due dates, subtasks, comments, files, activity, review, approval, and return-for-changes workflow.","Personal dashboard views: deadlines, calendar, notifications, profile skills, and a personal Softboard for sticky-note to-dos.","Project Brief upload/download so all members can access shared project instructions."]: bullet(doc,x)

h(doc,"2. Product Flow: From Sign-Up to Completed Work")
for x in ["A user registers. The backend hashes the password, creates a verification token, stores only its SHA-256 hash, and sends the raw token in the verification email.","The user verifies the email, then logs in. The API returns an access JWT and sets an HTTP-only refresh-token cookie.","A creator makes a project and becomes its Admin. They add existing users or invite new email addresses.","An Admin or Project Admin creates a task, assigns it, and may attach files. A member works on it, comments, uploads evidence, and submits it for review.","The reviewer approves the task to Completed or sends it back to In Progress. Activity entries and targeted notifications preserve visibility.","The frontend polls/refetches dashboard data periodically, while Softboard slot movement is optimistic: it updates visually first and saves in the background."]: numbered(doc,x,"product_flow")

h(doc,"3. Architecture at a Glance")
table(doc,["Layer","Technology","Responsibility"],[
["Browser client","React 19 + Vite + React Router","Screens, navigation, stateful forms, API calls, role-aware controls, themes, Softboard interaction."],
["API","Node.js + Express 5","REST endpoints, validation, authentication, authorization, business rules, JSON errors."],
["Data","MongoDB Atlas + Mongoose","Users, project memberships, tasks, comments, notes, activity, notifications, personal Softboard notes."],
["File storage","AWS S3 + AWS SDK + Multer","Private attachment and Project Brief storage; short-lived signed download URLs."],
["Email","SMTP / Mailtrap configuration","Verification, invitation, forgot-password email delivery."],
["Hosting","Vercel: separate frontend and API projects","Build, production deploy, live domain alias, environment-variable configuration."],
["Source control","GitHub + pull request workflow","Version history, reviewable commits, protected secrets through ignored environment files."]
],[1.2,1.7,3.6])
callout(doc,"Request path","React calls an Express REST route. The access JWT is sent as Bearer authorization. Express verifies the JWT, loads the user, verifies project membership/role where required, queries MongoDB, and returns a consistent JSON response. S3 and SMTP are used only when a file or email flow requires them.")

h(doc,"4. Technology Decisions: What, Why, and Alternatives")
table(doc,["Decision","Why it fits this project","Why not the common alternative?"],[
["React + Vite","Fast component development, client routing, familiar fresher stack, rapid local feedback.","Next.js was unnecessary because this authenticated dashboard does not need SEO/server-rendered public pages."],
["Express REST API","Clear middleware pipeline for JWT, validation, roles, uploads, and error handling.","GraphQL would add schema/resolver complexity without a strong need for client-selected nested data."],
["MongoDB + Mongoose","Natural fit for flexible task metadata and embedded arrays such as attachments/subtasks; Mongoose gives schema validation and model methods.","PostgreSQL is excellent for strict relational reporting, but MongoDB kept iteration fast for this placement project."],
["JWT access + refresh tokens","Scales naturally across a separated frontend/API; access token is sent per request and refresh cookie renews sessions.","A server-only session store is simpler for revocation but adds shared session-state infrastructure."],
["AWS S3 for files","Keeps binary files out of MongoDB and lets the API issue expiring private links.","Storing files as MongoDB blobs increases database backup/transfer burden for this use case."],
["Vercel","Simple Git/CLI deployment for React and serverless API projects.","Docker/Kubernetes are useful at larger scale but would add operations work without improving the placement demo."],
["REST polling","Enough for periodic dashboard freshness and easy to explain.","WebSockets/SSE are better for true real-time collaboration, but require connection lifecycle and scaling design."]
],[1.35,2.55,2.6])

h(doc,"5. Authentication and JWT — Explain This Correctly")
h(doc,"JWT in one minute",2)
add_text(doc,"A JSON Web Token is a signed compact token containing claims. It is not encrypted by default. The server signs it; later it verifies the signature before trusting the claims. In Task Force Orbit, the access token contains the user id, email, and username. The API verifies it with the access-token secret and then loads the user from MongoDB before authorizing the request.")
table(doc,["Token","How this project uses it","Security purpose"],[
["Access token","Returned in login response; frontend keeps it in JavaScript memory and sends Authorization: Bearer <token>.","Short-lived proof for protected API requests. Memory storage avoids persisting it in localStorage."],
["Refresh token","Stored on the user record and sent as an HTTP-only cookie scoped to /api/v1/auth; refresh endpoint rotates it and returns a new access token.","Longer-lived session renewal without exposing the refresh token to frontend JavaScript."],
["Verification / reset token","A random raw token is emailed; only its SHA-256 hash and expiry are stored.","If the database leaks, the raw email link token is not directly present there."]
],[1.2,2.8,2.5])
h(doc,"Actual login and refresh flow",2)
for x in ["Login verifies the password with bcrypt and refuses unverified accounts.","The API signs access and refresh JWTs, saves the refresh token on the user, returns the access token, and sets the refresh cookie as HTTP-only. In production it is Secure and SameSite=None because frontend/API are cross-origin.","The API client retries one failed 401 request by POSTing to /auth/refresh-token; if successful it stores the new access token only in memory and retries the original call.","Logout clears the stored refresh token and clears the cookie. This invalidates the currently stored refresh token server-side."]: numbered(doc,x,"auth_flow")
callout(doc,"Important interview nuance","Do not say JWT automatically means 'fully stateless'. The access token is self-contained, but this project deliberately stores the refresh token in MongoDB so logout and rotation can invalidate it. OWASP warns against storing authentication tokens in localStorage; this project keeps the refresh token HTTP-only and the access token in memory.")

h(doc,"6. Authorization, Roles, and Task State Machine")
add_text(doc,"Authentication answers 'who are you?'; authorization answers 'are you allowed to do this operation?'. Task Force Orbit performs both on the API, not only in React.")
table(doc,["Role","Typical capabilities","Backend enforcement"],[
["Admin","Create/delete project, manage members, manage tasks, review work, upload/update Project Brief.","verifyJWT + validateProjectPermission on role-restricted routes."],
["Project Admin","Manage and review work within a project.","Same server-side role middleware; UI controls are only a convenience."],
["Member","View project work, work on assigned tasks, comment, submit assigned work for review.","Submission controller verifies assignment; members cannot create/administer project tasks."]
],[1.25,2.65,2.6])
add_text(doc,"Task statuses are To Do, In Progress, In Review, and Completed. The review flow avoids putting a task back into To Do merely because it is under review: a reviewer either approves it to Completed or sends it to In Progress for changes. This is a stronger business rule than a generic drag-and-drop board.")

h(doc,"7. Data Modeling: MongoDB and Mongoose")
add_text(doc,"Core collections include User, Project, ProjectMember, Task, Comment, Note, Activity, Notification, Invitation, and SoftboardNote. Cross-entity relationships use ObjectId references: for example, a Task references its Project and assigned user. Small data naturally owned by one record is embedded: task attachments/subtasks and the Softboard note task checklist are arrays within their parent document.")
table(doc,["Modeling choice","Example here","Reason"],[
["Reference","ProjectMember links user and project; Task links project/assignee.","Roles and tasks are independently queried and grow separately."],
["Embed","Task attachments/subtasks; note checklist items.","The data is read/updated with the parent and remains bounded."],
["Index/constraint","Unique email and username; owner index on Softboard notes.","Protect identity uniqueness and make user-scoped retrieval efficient."]
],[1.25,2.6,2.65])
callout(doc,"Freshers follow-up","MongoDB writes are atomic at a single-document level. If a future operation needs all-or-nothing changes across several collections, discuss transactions, idempotency, or a compensating workflow rather than claiming the current code is automatically transactional everywhere.")

h(doc,"8. Files, Email, Notifications, and Softboard")
h(doc,"Private files",2)
add_text(doc,"Multer temporarily keeps accepted uploads in memory. Task attachments are limited to 1 MB and allowed MIME types; Project Briefs are limited to 5 MB and PDF/DOCX. The API uploads file buffers to private S3. MongoDB retains metadata and object key, not file bytes. For a download, the API generates a signed S3 GET URL valid for five minutes. Replacing/deleting supported work removes prior objects where the controller calls S3 cleanup.")
h(doc,"Email and notifications",2)
add_text(doc,"SMTP configuration sends verification, password-reset, and invitation messages. Database notifications are created for task events and reminders; the UI displays them and marks them read. This is in-app notification state, not a replacement for an enterprise push-notification service.")
h(doc,"Personal Softboard",2)
add_text(doc,"Softboard is user-specific, not project-specific. Notes belong to an owner, have a color, title, task checklist, and a slot index. It uses a fixed nine-slot board so placement is consistent across screen sizes. Movement is optimistic: React updates the selected slot immediately, then a queued API save persists it; this removes network latency from the interaction. Notes can be added up to nine and deleted by dragging to the bin.")

h(doc,"9. Security and Reliability Controls")
for x in ["bcrypt hashes passwords with cost factor 10; raw passwords are never returned by user queries.","Email verification and reset tokens are cryptographically random, hashed before database storage, and expire.","CORS accepts configured origins; a trusted-origin middleware rejects unsafe state-changing browser requests from untrusted origins. CORS is a browser policy, not authorization by itself.","Access tokens are validated on every protected API call; project membership and roles are validated before sensitive operations.","express-validator validates and sanitizes request bodies using matchedData; ObjectId parameters are validated; API errors are normalized to JSON.","Production API rate limiting is 100 requests/IP per 15 minutes. It is a basic abuse control, not DDoS protection.","Refresh cookies are HTTP-only and Secure in production. Secrets stay in Vercel environment variables and local .env files, not Git.","Private S3 objects are exposed through expiring signed URLs rather than public bucket URLs."]: bullet(doc,x)
callout(doc,"Honest limitations","For a stronger enterprise version: use a managed secrets/IAM-role approach instead of long-lived AWS keys; add audit/event monitoring, CSP/CSRF strategy tailored to the cookie design, email provider delivery monitoring, queues/retries, and a real-time transport for immediate multi-user updates.")

h(doc,"10. Development → Production → Live Updates")
table(doc,["Stage","What happens","How to explain it"],[
["Local development","React/Vite frontend runs locally; Express API runs locally; a separate test database is enforced.","Fast feedback without touching production data. Environment values are local and ignored by Git."],
["Automated verification","Node API tests cover health, auth, role boundaries, workflow, validation, rate limits, and database safety. Frontend has lint/tests/build.","Tests reduce regressions; they do not prove every UX path."],
["GitHub","Changes are intentionally staged, committed, pushed to codex/complete-backend-api, and visible in draft PR #21.","Small, descriptive commits make changes reviewable and reversible."],
["Vercel production","Frontend and API are separate Vercel projects. Production aliases serve the public site and API. Secrets are configured in the platform, not the repository.","A deployment builds an immutable release. The frontend production URL is https://www.taskforceorbit.space."],
["After deployment","Check frontend route response, API health endpoint, authenticated API behavior, and manually test critical flows such as login, project/task review, email, and files.","Deployment 'success' only means the build ran; functional verification is still required."]
],[1.35,2.85,2.3])
h(doc,"How current live updates work",2)
add_text(doc,"Normal dashboard freshness uses periodic refetching rather than WebSockets. Softboard interaction uses optimistic UI state: a note moves immediately on the client; saves are queued so quick moves persist in order. Another user will not see an update instantly unless they refetch/poll. A future WebSocket/SSE layer could publish events, but it would add authentication, reconnection, ordering, and scaling concerns.")

h(doc,"11. Placement Interview Question Bank")
questions=[
("Why is this more than a CRUD project?","It models a real review workflow: assign → work → attach/comment → submit for review → approve or return. It also includes permissions, audit/activity, notifications, files, email verification, deployment, and tests."),
("Why React and Express?","They separate UI concerns from API/business rules. React gives componentized screens; Express middleware makes authentication, validation, and authorization explicit and testable."),
("What is JWT?","A signed token with claims. The server verifies its signature before accepting it. It is not encrypted by default, so do not put secrets in its payload."),
("Access vs refresh token?","Access token authorizes requests and is short-lived; refresh token obtains a new access token. Here, access lives in memory and refresh is an HTTP-only cookie plus a server-stored value."),
("Why store refresh token in DB?","It enables logout/invalidation and rotation. The trade-off is server-side state for refresh sessions."),
("How do you prevent a Member from calling an admin API manually?","The Express route verifies JWT and then project membership/role. Hiding a frontend button alone is never trusted."),
("Why MongoDB?","Flexible task metadata and quick iteration. I still use Mongoose schemas, references, uniqueness, and validation rather than treating it as schema-free."),
("Why not store attachments in MongoDB?","S3 is better for file objects. MongoDB holds metadata/key; private signed links provide temporary access."),
("What is a presigned URL?","A time-limited URL signed by AWS credentials that grants access to one object operation without making the bucket public."),
("How do you secure credentials?","Secrets are environment variables in Vercel/local .env, excluded from Git. The code never sends AWS secrets to the browser."),
("What does CORS solve?","It is a browser cross-origin rule. It permits configured origins but does not replace server-side authentication/authorization."),
("How do verification links stay safer if DB leaks?","The server hashes the random token before storage; the raw value exists in the email link and is compared by hash."),
("How does the frontend recover from an expired access token?","The API helper receives 401, calls the refresh endpoint using the HTTP-only cookie, stores the returned access token in memory, and retries once."),
("How did you avoid Softboard drag lag?","Fixed slots avoid free-position layout work. On drop I update React state immediately, animate left/top briefly, and queue the database write in the background."),
("How would you make task updates real-time?","Add WebSockets or SSE, authenticate the connection, publish project events after durable writes, and refetch/merge safely on reconnect. I chose polling first to keep scope and operations simple."),
("What would you improve before a large production launch?","Observability, queue-based email, IAM roles/secrets manager, object lifecycle rules, stronger CSP/CSRF design, load testing, real-time events, and a documented backup/restore plan."),
("How do you demonstrate ownership?","Explain one bug or UX issue, the cause, the smallest safe fix, test evidence, Git commit, deployment, and live verification. The Softboard latency fix is a good example.")]
for q,a in questions:
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2); font(p.add_run("Q. "+q),10.8,True,BLUE)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); font(p.add_run("A. "+a),10.5)

h(doc,"12. Final Demo Script and Pre-Interview Checklist")
h(doc,"Five-minute demo order",2)
for x in ["State the problem and 60-second pitch.","Log in and point out verified authentication, dashboard, and role-aware navigation.","Open a project: create/describe task, show member/admin distinction, then show a task moving through In Progress → In Review → Completed.","Show comment/evidence upload and explain private S3 + five-minute signed link.","Open Project Brief and notifications/activity.","Open Softboard: show personal ownership, fixed slots, instant optimistic snap, and background persistence.","Close with deployment URL, GitHub PR/commit workflow, automated checks, and one honest future improvement."]: numbered(doc,x,"demo_flow")
h(doc,"Before the interview",2)
for x in ["Open the live URL and verify login with a safe demo account.","Know the exact project problem, one technical decision, one bug fix, one trade-off, and one future improvement.","Never reveal passwords, API keys, MongoDB URI, AWS credentials, or real user data.","If asked something not implemented, say so directly, explain how you would design it, and identify the trade-off.","Use business language first, then technical detail: problem → design → implementation → validation → result."]: bullet(doc,x)

h(doc,"13. Short Research Notes and Sources")
add_text(doc,"Focused research supports the interview emphasis in this guide: project discussion, reasoning about trade-offs, security fundamentals, and clear communication. The sources below are reference material; the implementation claims above were checked against the TaskForce repository.")
sources=[
"OWASP Session Management Cheat Sheet — https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html",
"OWASP Authentication Cheat Sheet — https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html",
"OWASP REST Security Cheat Sheet — https://cheatsheetseries.owasp.org/cheatsheets/REST_Security_Cheat_Sheet.html",
"MDN: Cross-Origin Resource Sharing (CORS) — https://developer.mozilla.org/en-US/docs/Web/HTTP/Guides/CORS",
"MongoDB: Embedded Data in Your Schema — https://www.mongodb.com/docs/manual/data-modeling/embedding/",
"MongoDB: Embedding vs References — https://www.mongodb.com/docs/v7.2/data-modeling/concepts/embedding-vs-references/",
"AWS S3: Download/upload objects with presigned URLs — https://docs.aws.amazon.com/AmazonS3/latest/userguide/using-presigned-url.html",
"Vercel secure SDLC guide (deployment/operations context) — https://assets.vercel.com/image/upload/v1745532177/front/security/Secured_modern_SDLC.pdf",
"Indeed: Entry-level software engineer interview questions (project/behavioral practice) — https://www.indeed.com/career-advice/interviewing/entry-level-software-engineer-interview-questions"
]
for s in sources: bullet(doc,s)
callout(doc,"Last reminder","Your goal is not to sound like you used every fashionable technology. Sound like an engineer who can explain a real product, protect user data, identify trade-offs, test changes, and improve the system safely.")

doc.save(OUT)
print(OUT)
